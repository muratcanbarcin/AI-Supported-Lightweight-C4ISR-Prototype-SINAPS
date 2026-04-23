"""
SINAPS C4ISR — Kapsamlı Test Senaryoları Excel Oluşturucu
===========================================================
Tüm modüller (Backend API, Frontend UI, AI Service, Security,
Simulation Engine, Integration) için detaylı test case'leri üretir.
"""

import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side,
    NamedStyle, numbers
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()

# ============================================================================
# STYLES
# ============================================================================
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
SECTION_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="1F4E79")
PASS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
SKIP_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

COLUMNS = [
    ("Test ID", 12),
    ("Modül", 18),
    ("Alt Modül", 22),
    ("Test Senaryosu", 50),
    ("Ön Koşullar", 35),
    ("Test Adımları", 55),
    ("Beklenen Sonuç", 40),
    ("Öncelik", 10),
    ("Sonuç", 12),
    ("Notlar / Hata Detayı", 30),
    ("Test Eden", 15),
    ("Tarih", 13),
]


def create_sheet(title: str, test_cases: list[dict]):
    """Create a formatted test sheet with the given test cases."""
    ws = wb.create_sheet(title=title)

    # Header row
    for col_idx, (col_name, col_width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    # Freeze header
    ws.freeze_panes = "A2"

    # Status dropdown validation
    status_dv = DataValidation(
        type="list",
        formula1='"BAŞARILI,BAŞARISIZ,ATLANMIŞ,TEST EDİLMEDİ"',
        allow_blank=True,
    )
    status_dv.error = "Lütfen geçerli bir sonuç seçin"
    status_dv.prompt = "Test sonucunu seçin"
    ws.add_data_validation(status_dv)

    priority_dv = DataValidation(
        type="list",
        formula1='"KRİTİK,YÜKSEK,ORTA,DÜŞÜK"',
        allow_blank=True,
    )
    ws.add_data_validation(priority_dv)

    row = 2
    current_section = None

    for tc in test_cases:
        # Section separator row
        section = tc.get("section")
        if section and section != current_section:
            current_section = section
            for col_idx in range(1, len(COLUMNS) + 1):
                cell = ws.cell(row=row, column=col_idx, value="")
                cell.fill = SECTION_FILL
                cell.border = THIN_BORDER
            ws.cell(row=row, column=1, value=section).font = SECTION_FONT
            ws.merge_cells(
                start_row=row, start_column=1,
                end_row=row, end_column=len(COLUMNS)
            )
            ws.cell(row=row, column=1).alignment = Alignment(horizontal="left", vertical="center")
            row += 1

        values = [
            tc.get("id", ""),
            tc.get("module", ""),
            tc.get("submodule", ""),
            tc.get("scenario", ""),
            tc.get("preconditions", ""),
            tc.get("steps", ""),
            tc.get("expected", ""),
            tc.get("priority", "ORTA"),
            "",  # Sonuç — boş (kullanıcı dolduracak)
            "",  # Notlar — boş
            "",  # Test Eden — boş
            "",  # Tarih — boş
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.alignment = WRAP
            cell.border = THIN_BORDER
            cell.font = Font(name="Calibri", size=10)

        # Apply dropdown to priority & status columns
        priority_dv.add(ws.cell(row=row, column=8))
        status_dv.add(ws.cell(row=row, column=9))

        row += 1

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{row - 1}"

    return ws


# ============================================================================
# TEST DATA — 1. AUTHENTICATION & SECURITY
# ============================================================================
auth_tests = [
    # --- Login ---
    {"section": "1. KULLANICI GİRİŞİ (LOGIN)", "id": "AUTH-001", "module": "Auth", "submodule": "Login",
     "scenario": "Geçerli kullanıcı adı ve şifre ile giriş",
     "preconditions": "- Veritabanında kayıtlı kullanıcı mevcut\n- Backend çalışıyor",
     "steps": "1. POST /api/v1/auth/login\n2. Body: {\"username\": \"admin\", \"password\": \"admin123\"}\n3. Response'u kontrol et",
     "expected": "200 OK\nJWT token, username ve role döner", "priority": "KRİTİK"},

    {"id": "AUTH-002", "module": "Auth", "submodule": "Login",
     "scenario": "Yanlış şifre ile giriş denemesi",
     "preconditions": "- Kayıtlı kullanıcı mevcut",
     "steps": "1. POST /api/v1/auth/login\n2. Body: {\"username\": \"admin\", \"password\": \"wrongpass\"}",
     "expected": "401 Unauthorized\n\"Invalid credentials\" mesajı", "priority": "KRİTİK"},

    {"id": "AUTH-003", "module": "Auth", "submodule": "Login",
     "scenario": "Var olmayan kullanıcı ile giriş",
     "preconditions": "- Backend çalışıyor",
     "steps": "1. POST /api/v1/auth/login\n2. Body: {\"username\": \"nonexistent\", \"password\": \"test\"}",
     "expected": "401 Unauthorized", "priority": "YÜKSEK"},

    {"id": "AUTH-004", "module": "Auth", "submodule": "Login",
     "scenario": "Boş username/password ile giriş",
     "preconditions": "- Backend çalışıyor",
     "steps": "1. POST /api/v1/auth/login\n2. Body: {\"username\": \"\", \"password\": \"\"}",
     "expected": "401 Unauthorized veya 400 Bad Request", "priority": "ORTA"},

    # --- JWT Token ---
    {"section": "2. JWT TOKEN YÖNETİMİ", "id": "AUTH-005", "module": "Auth", "submodule": "JWT",
     "scenario": "JWT token ile korumalı endpoint'e erişim",
     "preconditions": "- Geçerli JWT token alınmış",
     "steps": "1. GET /api/v1/admin/users\n2. Header: Authorization: Bearer <token>\n3. Response'u kontrol et",
     "expected": "200 OK\nKullanıcı listesi döner", "priority": "KRİTİK"},

    {"id": "AUTH-006", "module": "Auth", "submodule": "JWT",
     "scenario": "Token olmadan korumalı endpoint'e erişim",
     "preconditions": "- Backend çalışıyor",
     "steps": "1. GET /api/v1/admin/users\n2. Authorization header göndermeden istek at",
     "expected": "401 Unauthorized veya 403 Forbidden", "priority": "KRİTİK"},

    {"id": "AUTH-007", "module": "Auth", "submodule": "JWT",
     "scenario": "Süresi dolmuş (expired) JWT ile erişim",
     "preconditions": "- Expired JWT token mevcut",
     "steps": "1. GET /api/v1/admin/users\n2. Header: Authorization: Bearer <expired_token>",
     "expected": "401 Unauthorized", "priority": "YÜKSEK"},

    {"id": "AUTH-008", "module": "Auth", "submodule": "JWT",
     "scenario": "Geçersiz (manipüle edilmiş) JWT ile erişim",
     "preconditions": "- Bozulmuş JWT token",
     "steps": "1. GET /api/v1/admin/users\n2. Header: Authorization: Bearer invalid.token.here",
     "expected": "401 Unauthorized", "priority": "YÜKSEK"},

    # --- RBAC ---
    {"section": "3. ROL TABANLI ERİŞİM KONTROLÜ (RBAC)", "id": "AUTH-009", "module": "Auth", "submodule": "RBAC",
     "scenario": "KARARGAH_ADMIN rolü ile admin endpoint'e erişim",
     "preconditions": "- KARARGAH_ADMIN rolündeki kullanıcı ile giriş yapılmış",
     "steps": "1. GET /api/v1/admin/users\n2. KARARGAH_ADMIN token ile istek at",
     "expected": "200 OK\nTüm kullanıcılar listelenir", "priority": "KRİTİK"},

    {"id": "AUTH-010", "module": "Auth", "submodule": "RBAC",
     "scenario": "HAREKAT_SUBAYI rolü ile admin endpoint'e erişim (yetkisiz)",
     "preconditions": "- HAREKAT_SUBAYI rolündeki kullanıcı ile giriş",
     "steps": "1. GET /api/v1/admin/users\n2. HAREKAT_SUBAYI token ile istek at",
     "expected": "403 Forbidden", "priority": "KRİTİK"},

    {"id": "AUTH-011", "module": "Auth", "submodule": "RBAC",
     "scenario": "HAREKAT_SUBAYI sadece kendi zone'undaki asset'leri görebilmeli",
     "preconditions": "- HAREKAT_SUBAYI kullanıcı assignedZoneId ile tanımlı\n- Birden fazla zone ve asset mevcut",
     "steps": "1. HAREKAT_SUBAYI token ile GET /api/v1/assets\n2. Dönen asset listesinin tamamının assignedZoneId'ye ait olduğunu doğrula",
     "expected": "Sadece atanmış zone'a ait asset'ler döner\nDiğer zone'ların asset'leri görünmez", "priority": "KRİTİK"},

    {"id": "AUTH-012", "module": "Auth", "submodule": "RBAC",
     "scenario": "ORDU_KOMUTANI rolü ile asset listeleme",
     "preconditions": "- ORDU_KOMUTANI rolündeki kullanıcı ile giriş",
     "steps": "1. GET /api/v1/assets\n2. ORDU_KOMUTANI token ile istek at",
     "expected": "Atanmış zone dahilindeki asset'ler döner", "priority": "YÜKSEK"},

    {"id": "AUTH-013", "module": "Auth", "submodule": "RBAC",
     "scenario": "HAREKAT_SUBAYI ile asset oluşturma denemesi (yetkisiz)",
     "preconditions": "- HAREKAT_SUBAYI token hazır",
     "steps": "1. POST /api/v1/assets\n2. HAREKAT_SUBAYI token ile asset oluşturma body gönder",
     "expected": "403 Forbidden\nSadece KARARGAH_ADMIN oluşturabilir", "priority": "YÜKSEK"},
]

# ============================================================================
# TEST DATA — 2. ASSET MANAGEMENT
# ============================================================================
asset_tests = [
    {"section": "4. ASSET LİSTELEME", "id": "AST-001", "module": "Asset", "submodule": "Listeleme",
     "scenario": "Tüm asset'leri listeleme (Admin)",
     "preconditions": "- KARARGAH_ADMIN ile giriş\n- DB'de asset kayıtları var",
     "steps": "1. GET /api/v1/assets\n2. Authorization header ekle\n3. Response kontrol et",
     "expected": "200 OK\nJSON array döner\nHer asset: id, name, type, inventoryStatus, isSimulated alanları içerir", "priority": "KRİTİK"},

    {"id": "AST-002", "module": "Asset", "submodule": "Listeleme",
     "scenario": "Boş veritabanında asset listeleme",
     "preconditions": "- DB'de asset yok\n- Admin girişi yapılmış",
     "steps": "1. GET /api/v1/assets",
     "expected": "200 OK\nBoş JSON array [] döner", "priority": "ORTA"},

    {"section": "5. ASSET OLUŞTURMA (CRUD)", "id": "AST-003", "module": "Asset", "submodule": "Oluşturma",
     "scenario": "Yeni asset oluşturma (tüm alanlar dolu)",
     "preconditions": "- KARARGAH_ADMIN ile giriş\n- Geçerli lat/lon değerleri",
     "steps": "1. POST /api/v1/assets\n2. Body: {\n  \"id\": \"TEST-DRONE-01\",\n  \"type\": \"DRONE\",\n  \"latitude\": 39.95,\n  \"longitude\": 32.85,\n  \"isSimulated\": true,\n  \"assignedZoneId\": \"<zone-uuid>\"\n}\n3. Response kontrol et",
     "expected": "200 OK\nOluşturulan asset JSON döner\nOtomatik telemetry kaydı oluşur", "priority": "KRİTİK"},

    {"id": "AST-004", "module": "Asset", "submodule": "Oluşturma",
     "scenario": "ID olmadan asset oluşturma denemesi",
     "preconditions": "- KARARGAH_ADMIN ile giriş",
     "steps": "1. POST /api/v1/assets\n2. Body: {\"type\": \"DRONE\", \"latitude\": 39.95, \"longitude\": 32.85}",
     "expected": "400 Bad Request\n\"id is required\" hata mesajı", "priority": "YÜKSEK"},

    {"id": "AST-005", "module": "Asset", "submodule": "Oluşturma",
     "scenario": "Koordinat olmadan asset oluşturma",
     "preconditions": "- KARARGAH_ADMIN ile giriş",
     "steps": "1. POST /api/v1/assets\n2. Body: {\"id\": \"TEST-001\", \"type\": \"VEHICLE\"}",
     "expected": "400 Bad Request\n\"latitude/longitude are required\" mesajı", "priority": "YÜKSEK"},

    {"id": "AST-006", "module": "Asset", "submodule": "Oluşturma",
     "scenario": "Aynı isimde asset oluşturma (duplicate)",
     "preconditions": "- Aynı name ile bir asset zaten mevcut",
     "steps": "1. POST /api/v1/assets\n2. Mevcut asset name ile body gönder",
     "expected": "400 Bad Request\n\"Asset id already exists\" mesajı", "priority": "YÜKSEK"},

    {"id": "AST-007", "module": "Asset", "submodule": "Oluşturma",
     "scenario": "Tüm asset tiplerini oluşturma (DRONE, VEHICLE, SOLDIER, SENSOR)",
     "preconditions": "- KARARGAH_ADMIN ile giriş",
     "steps": "1. Her tip için ayrı POST /api/v1/assets isteği gönder\n2. Her birinin başarılı oluştuğunu doğrula",
     "expected": "Her tip için 200 OK\nTip alanı doğru kaydedilir", "priority": "ORTA"},

    {"section": "6. ASSET SİLME", "id": "AST-008", "module": "Asset", "submodule": "Silme",
     "scenario": "Mevcut asset'i silme",
     "preconditions": "- KARARGAH_ADMIN ile giriş\n- Silinecek asset mevcut",
     "steps": "1. DELETE /api/v1/assets/{id}\n2. GET ile asset'in silindiğini doğrula",
     "expected": "200 OK\nAsset ve ilişkili telemetry/task kayıtları silinir", "priority": "KRİTİK"},

    {"id": "AST-009", "module": "Asset", "submodule": "Silme",
     "scenario": "Var olmayan asset'i silme",
     "preconditions": "- KARARGAH_ADMIN ile giriş",
     "steps": "1. DELETE /api/v1/assets/NON-EXISTENT-ID",
     "expected": "400 Bad Request\n\"Asset not found\" mesajı", "priority": "ORTA"},

    {"section": "7. ACİL DURDURMA (EMERGENCY STOP)", "id": "AST-010", "module": "Asset", "submodule": "Emergency Stop",
     "scenario": "Asset'e acil durdurma komutu gönderme",
     "preconditions": "- Aktif telemetry verisi olan bir asset mevcut",
     "steps": "1. POST /api/v1/assets/{id}/emergency-stop\n2. Response kontrol et\n3. Asset'in son telemetry payload'ını kontrol et",
     "expected": "200 OK\n{\"success\": true}\nTelemetry payload'da status: STOPPED\nsystem_logs tablosunda EMERGENCY_STOP kaydı", "priority": "KRİTİK"},

    {"id": "AST-011", "module": "Asset", "submodule": "Emergency Stop",
     "scenario": "Telemetrisi olmayan asset'e emergency stop",
     "preconditions": "- Telemetrisi olmayan asset mevcut",
     "steps": "1. POST /api/v1/assets/{id}/emergency-stop",
     "expected": "200 OK\nBaşarılı (sadece system_logs kaydı oluşur)", "priority": "ORTA"},
]

# ============================================================================
# TEST DATA — 3. TELEMETRY
# ============================================================================
telemetry_tests = [
    {"section": "8. TELEMETRİ SORGULAMA", "id": "TLM-001", "module": "Telemetry", "submodule": "Latest",
     "scenario": "En güncel telemetri verilerini çekme",
     "preconditions": "- Simülasyon çalışıyor\n- Telemetry tablosunda veri var",
     "steps": "1. GET /api/v1/telemetry/latest\n2. Response JSON array kontrol et",
     "expected": "200 OK\nHer asset için en son telemetry kaydı\nlatitude, longitude, speed, batteryLevel, sourceId, sourceType alanları", "priority": "KRİTİK"},

    {"id": "TLM-002", "module": "Telemetry", "submodule": "History",
     "scenario": "Telemetri geçmişini çekme",
     "preconditions": "- Telemetry tablosunda geçmiş veriler var",
     "steps": "1. GET /api/v1/telemetry/history\n2. Response kontrol et",
     "expected": "200 OK\nTüm geçmiş telemetry kayıtları döner", "priority": "YÜKSEK"},

    {"section": "9. TELEMETRİ EKLEME", "id": "TLM-003", "module": "Telemetry", "submodule": "Ingest",
     "scenario": "Tek telemetri verisi gönderme (POST)",
     "preconditions": "- Backend çalışıyor",
     "steps": "1. POST /api/v1/telemetry\n2. Body: {\"domain\":\"drone\", \"type\":\"telemetry\", \"id\":\"DRONE-TEST\", \"time\":\"...\", \"payload\":{\"position\":{\"lat\":39.95,\"lon\":32.60,\"alt_m\":150},\"ground_speed_mps\":25.5,\"battery\":{\"soc_pct\":85}}}",
     "expected": "200 OK\n{\"status\": \"OK\"}\nVeri telemetry tablosuna kaydedilir", "priority": "KRİTİK"},

    {"id": "TLM-004", "module": "Telemetry", "submodule": "Ingest",
     "scenario": "Batch telemetri gönderme",
     "preconditions": "- Backend çalışıyor",
     "steps": "1. POST /api/v1/telemetry/batch\n2. Body: JSON array (birden fazla telemetri kaydı)",
     "expected": "200 OK\n\"Processed N telemetry records\" mesajı", "priority": "YÜKSEK"},

    {"id": "TLM-005", "module": "Telemetry", "submodule": "Ingest",
     "scenario": "Geçersiz format ile telemetri gönderme",
     "preconditions": "- Backend çalışıyor",
     "steps": "1. POST /api/v1/telemetry\n2. Body: {\"invalid\": \"data\"}",
     "expected": "400 Bad Request\nHata mesajı döner", "priority": "ORTA"},

    {"section": "10. TELEMETRİ HARİTA POLLING", "id": "TLM-006", "module": "Telemetry", "submodule": "Polling",
     "scenario": "3 saniyede bir /telemetry/latest pollingi",
     "preconditions": "- Frontend harita sayfası açık\n- Simülasyon çalışıyor",
     "steps": "1. Network tab'ında /telemetry/latest isteklerini izle\n2. Her 3 saniyede bir istek atıldığını doğrula\n3. Konum verilerinin güncellendiğini kontrol et",
     "expected": "Düzenli aralıklarla veri çekiliyor\nHarita üzerinde birimler güncelleniyor", "priority": "KRİTİK"},
]

# ============================================================================
# TEST DATA — 4. TACTICAL ZONES
# ============================================================================
zone_tests = [
    {"section": "11. TAKTİK BÖLGE YÖNETİMİ", "id": "ZON-001", "module": "Zones", "submodule": "Listeleme",
     "scenario": "Tüm taktik bölgeleri listeleme",
     "preconditions": "- DB'de zone kayıtları var",
     "steps": "1. GET /api/v1/zones\n2. Response kontrol et",
     "expected": "200 OK\nHer zone: id (UUID), name, zoneType, latitude, longitude, radiusMeters", "priority": "KRİTİK"},

    {"id": "ZON-002", "module": "Zones", "submodule": "Oluşturma",
     "scenario": "Yeni taktik bölge oluşturma",
     "preconditions": "- KARARGAH_ADMIN ile giriş",
     "steps": "1. POST /api/v1/zones\n2. Body: {\"name\":\"Test Bölgesi\", \"zoneType\":\"AO\", \"latitude\":39.95, \"longitude\":32.85, \"radiusMeters\":5000}",
     "expected": "200 OK\nOluşturulan zone UUID ile döner", "priority": "KRİTİK"},

    {"id": "ZON-003", "module": "Zones", "submodule": "Güncelleme",
     "scenario": "Taktik bölge güncelleme",
     "preconditions": "- KARARGAH_ADMIN ile giriş\n- Mevcut zone var",
     "steps": "1. PUT /api/v1/zones/{id}\n2. Güncellenmiş body gönder",
     "expected": "200 OK\nGüncellenmiş zone döner", "priority": "YÜKSEK"},

    {"id": "ZON-004", "module": "Zones", "submodule": "Silme",
     "scenario": "Taktik bölge silme",
     "preconditions": "- KARARGAH_ADMIN ile giriş\n- Silinecek zone mevcut",
     "steps": "1. DELETE /api/v1/zones/{id}\n2. GET ile silindiğini doğrula",
     "expected": "200 OK\nZone silinir", "priority": "YÜKSEK"},

    {"id": "ZON-005", "module": "Zones", "submodule": "RBAC",
     "scenario": "HAREKAT_SUBAYI ile zone oluşturma denemesi",
     "preconditions": "- HAREKAT_SUBAYI ile giriş",
     "steps": "1. POST /api/v1/zones\n2. HAREKAT_SUBAYI token ile gönder",
     "expected": "403 Forbidden", "priority": "YÜKSEK"},
]

# ============================================================================
# TEST DATA — 5. TASKS & DEPLOYMENT
# ============================================================================
task_tests = [
    {"section": "12. GÖREV YÖNETİMİ (TASKS)", "id": "TSK-001", "module": "Tasks", "submodule": "Listeleme",
     "scenario": "Tüm görevleri listeleme",
     "preconditions": "- Task kayıtları mevcut",
     "steps": "1. GET /api/v1/tasks",
     "expected": "200 OK\nGörev listesi (id, assetId, title, status, taskType, zoneName)", "priority": "KRİTİK"},

    {"id": "TSK-002", "module": "Tasks", "submodule": "Listeleme",
     "scenario": "Asset ID'ye göre görev filtreleme",
     "preconditions": "- Belirli asset'e ait görevler var",
     "steps": "1. GET /api/v1/tasks?assetId=DRONE-001",
     "expected": "200 OK\nSadece ilgili asset'in görevleri döner", "priority": "YÜKSEK"},

    {"id": "TSK-003", "module": "Tasks", "submodule": "Oluşturma",
     "scenario": "Yeni görev oluşturma",
     "preconditions": "- Backend çalışıyor",
     "steps": "1. POST /api/v1/tasks\n2. Body: {\"assetId\":\"DRONE-001\", \"title\":\"Keşif Görevi\", \"description\":\"Bölge keşfi\"}",
     "expected": "200 OK\nStatus otomatik PENDING olur", "priority": "KRİTİK"},

    {"section": "13. KONUŞLANDIRMA (DEPLOY)", "id": "TSK-004", "module": "Tasks", "submodule": "Deploy",
     "scenario": "Asset'i taktik bölgeye konuşlandırma",
     "preconditions": "- Mevcut asset ve zone kaydı var",
     "steps": "1. POST /api/v1/tasks/deploy\n2. Body: {\"assetId\":\"DRONE-001\", \"zoneId\":\"<zone-uuid>\"}",
     "expected": "200 OK\nDeploy task oluşur\nAsset zone'a doğru hareket etmeye başlar", "priority": "KRİTİK"},

    {"id": "TSK-005", "module": "Tasks", "submodule": "Deploy",
     "scenario": "Geçersiz zone ID ile deploy denemesi",
     "preconditions": "- Mevcut asset var\n- Geçersiz zone UUID",
     "steps": "1. POST /api/v1/tasks/deploy\n2. Body: {\"assetId\":\"DRONE-001\", \"zoneId\":\"invalid-uuid\"}",
     "expected": "400 Bad Request\nHata mesajı döner", "priority": "YÜKSEK"},

    {"id": "TSK-006", "module": "Tasks", "submodule": "Deploy",
     "scenario": "Var olmayan asset ile deploy denemesi",
     "preconditions": "- Geçerli zone mevcut",
     "steps": "1. POST /api/v1/tasks/deploy\n2. Body: {\"assetId\":\"NONEXISTENT\", \"zoneId\":\"<valid-uuid>\"}",
     "expected": "400 Bad Request\n\"Asset not found\" mesajı", "priority": "YÜKSEK"},
]

# ============================================================================
# TEST DATA — 6. ALERTS & ANALYSIS
# ============================================================================
alert_tests = [
    {"section": "14. ALARM YÖNETİMİ", "id": "ALR-001", "module": "Alerts", "submodule": "Listeleme",
     "scenario": "Tüm alarmları listeleme",
     "preconditions": "- Alarm kayıtları mevcut",
     "steps": "1. GET /api/v1/alerts",
     "expected": "200 OK\nAlarm listesi (id, type, sourceId, severity, message, acknowledged)", "priority": "KRİTİK"},

    {"id": "ALR-002", "module": "Alerts", "submodule": "Listeleme",
     "scenario": "Onaylanmamış alarmları listeleme",
     "preconditions": "- Onaylanmamış alarm var",
     "steps": "1. GET /api/v1/alerts/unacknowledged",
     "expected": "200 OK\nSadece acknowledged=false olanlar döner", "priority": "YÜKSEK"},

    {"id": "ALR-003", "module": "Alerts", "submodule": "Filtreleme",
     "scenario": "Asset'e göre alarm filtreleme",
     "preconditions": "- Belirli asset'e ait alarmlar var",
     "steps": "1. GET /api/v1/alerts/asset/{sourceId}",
     "expected": "200 OK\nSadece ilgili asset'in alarmları", "priority": "YÜKSEK"},

    {"id": "ALR-004", "module": "Alerts", "submodule": "İstatistik",
     "scenario": "Alarm istatistiklerini görüntüleme",
     "preconditions": "- Alarm kayıtları var",
     "steps": "1. GET /api/v1/alerts/stats",
     "expected": "200 OK\nToplam alarm, onaylanan, kritik sayısı", "priority": "ORTA"},

    {"id": "ALR-005", "module": "Alerts", "submodule": "Onaylama",
     "scenario": "Alarm onaylama (acknowledge)",
     "preconditions": "- Onaylanmamış alarm var",
     "steps": "1. POST /api/v1/alerts/{id}/acknowledge",
     "expected": "200 OK\n{\"success\": true}\nAlarm acknowledged=true olur", "priority": "KRİTİK"},

    {"id": "ALR-006", "module": "Alerts", "submodule": "Silme",
     "scenario": "Alarm silme",
     "preconditions": "- Mevcut alarm var",
     "steps": "1. DELETE /api/v1/alerts/{id}",
     "expected": "200 OK\n{\"success\": true}", "priority": "YÜKSEK"},

    {"section": "15. TAKTİK ANALİZ", "id": "ANL-001", "module": "Analysis", "submodule": "Listeleme",
     "scenario": "Tüm taktik analizleri listeleme",
     "preconditions": "- Tactical analysis kayıtları var",
     "steps": "1. GET /api/v1/analysis",
     "expected": "200 OK\nthreatLevel, decision, reasoning, confidence alanları", "priority": "KRİTİK"},

    {"id": "ANL-002", "module": "Analysis", "submodule": "Filtreleme",
     "scenario": "Tehdit seviyesine göre filtreleme",
     "preconditions": "- Farklı threat level'larda analiz kayıtları var",
     "steps": "1. GET /api/v1/analysis/threat/HIGH",
     "expected": "200 OK\nSadece HIGH threat level kayıtlar döner", "priority": "YÜKSEK"},

    {"id": "ANL-003", "module": "Analysis", "submodule": "Detay",
     "scenario": "Tekil analiz detayı görüntüleme",
     "preconditions": "- Analiz kaydı mevcut",
     "steps": "1. GET /api/v1/analysis/{id}",
     "expected": "200 OK\nTek analiz detayı (threatLevel, decision, telemetry bilgisi)", "priority": "ORTA"},
]

# ============================================================================
# TEST DATA — 7. AI SERVICE
# ============================================================================
ai_tests = [
    {"section": "16. AI CHAT (Llama 3.1 8B)", "id": "AI-001", "module": "AI Service", "submodule": "Chat",
     "scenario": "Türkçe soru ile AI chat",
     "preconditions": "- AI Service çalışıyor (localhost:8000)\n- Llama modeli yüklü",
     "steps": "1. POST /api/v1/ai/chat\n2. Body: {\"query\": \"Drone bataryası %15 altına düştüğünde ne yapılmalı?\"}",
     "expected": "200 OK\nTürkçe yanıt\ncited_rule (varsa)\nrag_hits sayısı\nprocessing_time_ms", "priority": "KRİTİK"},

    {"id": "AI-002", "module": "AI Service", "submodule": "Chat",
     "scenario": "Boş sorgu ile AI chat",
     "preconditions": "- AI Service çalışıyor",
     "steps": "1. POST /api/v1/ai/chat\n2. Body: {\"query\": \"\"}",
     "expected": "400 Bad Request\n\"Sorgu boş olamaz\" mesajı", "priority": "YÜKSEK"},

    {"id": "AI-003", "module": "AI Service", "submodule": "Chat",
     "scenario": "@Mention ile asset hakkında soru sorma",
     "preconditions": "- AI Service çalışıyor\n- DB'de ilgili asset mevcut",
     "steps": "1. POST /api/v1/ai/chat\n2. Body: {\"query\": \"@2ORDU-TB2-01 bu droneun batarya durumu nedir?\"}",
     "expected": "200 OK\nYanıtta asset'in gerçek zamanlı verileri kullanılır\nKonum, batarya, hız bilgileri doğru", "priority": "KRİTİK"},

    {"id": "AI-004", "module": "AI Service", "submodule": "Chat",
     "scenario": "RAG retrieval doğrulama (doktrin eşleşmesi)",
     "preconditions": "- tactical_memory tablosunda doktrin kuralları var",
     "steps": "1. POST /api/v1/ai/chat\n2. IHA ile ilgili soru sor\n3. cited_rule alanını kontrol et",
     "expected": "cited_rule IHA/drone doktrin kuralıyla eşleşir\nAraç kuralı dönmez (cross-domain validation)", "priority": "YÜKSEK"},

    {"section": "17. AI ANALİZ (Gemma 2 9B)", "id": "AI-005", "module": "AI Service", "submodule": "Analyze",
     "scenario": "Drone telemetrisi ile tehdit analizi",
     "preconditions": "- AI Service çalışıyor\n- Gemma modeli yüklü",
     "steps": "1. POST localhost:8000/analyze\n2. Drone telemetry JSON body\n3. Response kontrol et",
     "expected": "decision, threat_level, confidence, reasoning alanları\nJSON formatında yapılandırılmış yanıt", "priority": "KRİTİK"},

    {"id": "AI-006", "module": "AI Service", "submodule": "Analyze",
     "scenario": "Asker (soldier) biyometrik analizi",
     "preconditions": "- AI Service çalışıyor",
     "steps": "1. POST localhost:8000/analyze\n2. Soldier telemetry (vitals: hr_bpm, spo2) body",
     "expected": "Biyometrik risk değerlendirmesi\nTehdit seviyesi ve öneriler", "priority": "YÜKSEK"},

    {"id": "AI-007", "module": "AI Service", "submodule": "Analyze",
     "scenario": "Araç telemetrisi analizi (düşük yakıt)",
     "preconditions": "- AI Service çalışıyor",
     "steps": "1. POST localhost:8000/analyze\n2. Vehicle telemetry (fuel_pct: 8) body",
     "expected": "Yüksek risk tespit edilir\nYakıt ikmal önerisi", "priority": "YÜKSEK"},

    {"section": "18. AI SERVİS SAĞLIK", "id": "AI-008", "module": "AI Service", "submodule": "Health",
     "scenario": "AI servis health check",
     "preconditions": "- AI Service çalışıyor",
     "steps": "1. GET localhost:8000/health",
     "expected": "200 OK\nstatus: healthy\nversion bilgisi", "priority": "KRİTİK"},

    {"id": "AI-009", "module": "AI Service", "submodule": "Info",
     "scenario": "AI servis bilgi endpoint'i",
     "preconditions": "- AI Service çalışıyor",
     "steps": "1. GET localhost:8000/info",
     "expected": "200 OK\nEngines bilgisi (Gemma + Llama)\nRAG model bilgisi\nDesteklenen tipler listesi", "priority": "ORTA"},

    {"section": "19. RAG MEMORY", "id": "AI-010", "module": "AI Service", "submodule": "Memory Ingest",
     "scenario": "Yeni angajman kuralı ekleme (ingest)",
     "preconditions": "- AI Service çalışıyor\n- pgvector bağlantısı aktif",
     "steps": "1. POST localhost:8000/ingest\n2. Body: {\"event_description\": \"IHA bataryası %20 altına düşerse...\", \"threat_category\": \"ENGAGEMENT_RULE\"}",
     "expected": "200 OK\nsuccess: true\nevent_id döner\nEmbedding oluşturulur", "priority": "YÜKSEK"},

    {"id": "AI-011", "module": "AI Service", "submodule": "Memory Stats",
     "scenario": "Bellek istatistiklerini görüntüleme",
     "preconditions": "- AI Service çalışıyor",
     "steps": "1. GET localhost:8000/memory/stats",
     "expected": "200 OK\nToplam kayıt sayısı\nKategori dağılımı", "priority": "ORTA"},
]

# ============================================================================
# TEST DATA — 8. CHAT SESSIONS
# ============================================================================
chat_tests = [
    {"section": "20. SOHBET OTURUMLARI", "id": "CHT-001", "module": "Chat", "submodule": "Session CRUD",
     "scenario": "Yeni sohbet oturumu oluşturma",
     "preconditions": "- Backend çalışıyor",
     "steps": "1. POST /api/v1/chats\n2. Response kontrol et",
     "expected": "200 OK\nid, title, createdAt alanları döner", "priority": "KRİTİK"},

    {"id": "CHT-002", "module": "Chat", "submodule": "Session CRUD",
     "scenario": "Tüm oturumları listeleme",
     "preconditions": "- Oturum kayıtları var",
     "steps": "1. GET /api/v1/chats",
     "expected": "200 OK\nSon güncellenene göre sıralı oturum listesi", "priority": "YÜKSEK"},

    {"id": "CHT-003", "module": "Chat", "submodule": "Session CRUD",
     "scenario": "Oturum silme",
     "preconditions": "- Mevcut oturum var",
     "steps": "1. DELETE /api/v1/chats/{sessionId}\n2. Oturumun ve mesajlarının silindiğini doğrula",
     "expected": "204 No Content\nOturum ve mesajlar silinir", "priority": "YÜKSEK"},

    {"section": "21. SOHBET MESAJLARI", "id": "CHT-004", "module": "Chat", "submodule": "Messages",
     "scenario": "Oturumda mesaj gönderme (user→AI round-trip)",
     "preconditions": "- Mevcut oturum var\n- AI Service çalışıyor",
     "steps": "1. POST /api/v1/chats/{sessionId}/messages\n2. Body: {\"content\": \"Hava savunma protokolü nedir?\"}",
     "expected": "200 OK\nAI yanıtı döner\nHer iki mesaj (user+assistant) DB'ye kaydedilir", "priority": "KRİTİK"},

    {"id": "CHT-005", "module": "Chat", "submodule": "Messages",
     "scenario": "Oturum mesajlarını çekme",
     "preconditions": "- Mesajları olan oturum mevcut",
     "steps": "1. GET /api/v1/chats/{sessionId}/messages",
     "expected": "200 OK\nKronolojik sırada mesaj listesi (user/assistant rolleri)", "priority": "YÜKSEK"},

    {"id": "CHT-006", "module": "Chat", "submodule": "Messages",
     "scenario": "Boş mesaj gönderme",
     "preconditions": "- Mevcut oturum var",
     "steps": "1. POST /api/v1/chats/{sessionId}/messages\n2. Body: {\"content\": \"\"}",
     "expected": "400 Bad Request\n\"Mesaj içeriği boş olamaz\" mesajı", "priority": "ORTA"},

    {"id": "CHT-007", "module": "Chat", "submodule": "Messages",
     "scenario": "İlk mesajda otomatik oturum başlığı oluşturma",
     "preconditions": "- Yeni boş oturum",
     "steps": "1. POST /api/v1/chats/{sessionId}/messages\n2. İlk mesajı gönder\n3. GET /api/v1/chats ile oturum başlığını kontrol et",
     "expected": "Oturum title'ı otomatik güncellenir\nMesaj içeriğinden türetilir", "priority": "ORTA"},
]

# ============================================================================
# TEST DATA — 9. ADMIN PANEL
# ============================================================================
admin_tests = [
    {"section": "22. KULLANICI YÖNETİMİ (ADMIN)", "id": "ADM-001", "module": "Admin", "submodule": "User CRUD",
     "scenario": "Yeni kullanıcı oluşturma",
     "preconditions": "- KARARGAH_ADMIN ile giriş",
     "steps": "1. POST /api/v1/admin/users\n2. Body: {\"username\":\"test_user\", \"password\":\"Test123!\", \"role\":\"HAREKAT_SUBAYI\", \"assignedZoneId\":\"<zone-uuid>\"}",
     "expected": "200 OK\nKullanıcı BCrypt hashli şifre ile kaydedilir\nassignedZoneId doğru atanır", "priority": "KRİTİK"},

    {"id": "ADM-002", "module": "Admin", "submodule": "User CRUD",
     "scenario": "Aynı username ile kullanıcı oluşturma (duplicate)",
     "preconditions": "- Aynı username'de kullanıcı zaten var",
     "steps": "1. POST /api/v1/admin/users\n2. Mevcut username ile body gönder",
     "expected": "400 Bad Request\n\"Username already exists\" mesajı", "priority": "YÜKSEK"},

    {"id": "ADM-003", "module": "Admin", "submodule": "User CRUD",
     "scenario": "HAREKAT_SUBAYI oluşturma (assignedZoneId olmadan)",
     "preconditions": "- KARARGAH_ADMIN ile giriş",
     "steps": "1. POST /api/v1/admin/users\n2. Body: {\"username\":\"subay1\", \"password\":\"pass\", \"role\":\"HAREKAT_SUBAYI\"}",
     "expected": "400 Bad Request\n\"assignedZoneId is required for HAREKAT_SUBAYI\" mesajı", "priority": "YÜKSEK"},

    {"id": "ADM-004", "module": "Admin", "submodule": "User CRUD",
     "scenario": "Kullanıcı silme",
     "preconditions": "- KARARGAH_ADMIN ile giriş\n- Silinecek kullanıcı var",
     "steps": "1. DELETE /api/v1/admin/users/{uuid}",
     "expected": "204 No Content\nKullanıcı silinir", "priority": "YÜKSEK"},

    {"section": "23. SİSTEM DURUMU", "id": "ADM-005", "module": "Admin", "submodule": "System Status",
     "scenario": "Sistem durumu sorgulama",
     "preconditions": "- KARARGAH_ADMIN ile giriş\n- Backend çalışıyor",
     "steps": "1. GET /api/v1/admin/system-status",
     "expected": "200 OK\nstatus: ONLINE\nactiveUsers, totalAssets, activeTasks, totalZones sayıları", "priority": "KRİTİK"},

    {"id": "ADM-006", "module": "Admin", "submodule": "Service Status",
     "scenario": "Servis sağlık durumu kontrolü",
     "preconditions": "- KARARGAH_ADMIN ile giriş",
     "steps": "1. GET /api/v1/admin/service-status",
     "expected": "200 OK\nbackend: ONLINE\ndatabase: ONLINE/OFFLINE\naiService: ONLINE/OFFLINE\nmqttBroker: ONLINE/OFFLINE\ndockerEngine: durumu", "priority": "KRİTİK"},
]

# ============================================================================
# TEST DATA — 10. SYSTEM LOGS
# ============================================================================
log_tests = [
    {"section": "24. SİSTEM LOGLARI", "id": "LOG-001", "module": "System Logs", "submodule": "Listeleme",
     "scenario": "Tüm logları listeleme",
     "preconditions": "- Log kayıtları mevcut",
     "steps": "1. GET /api/v1/system-logs",
     "expected": "200 OK\nSon oluşturulanlar önce\nassetId, eventType, message alanları", "priority": "YÜKSEK"},

    {"id": "LOG-002", "module": "System Logs", "submodule": "Filtreleme",
     "scenario": "Asset'e göre log filtreleme",
     "preconditions": "- Belirli asset'e ait loglar var",
     "steps": "1. GET /api/v1/system-logs/{assetId}",
     "expected": "200 OK\nSadece ilgili asset'in logları", "priority": "YÜKSEK"},

    {"id": "LOG-003", "module": "System Logs", "submodule": "Son Loglar",
     "scenario": "Son N logu getirme",
     "preconditions": "- Çok sayıda log var",
     "steps": "1. GET /api/v1/system-logs/recent?limit=10",
     "expected": "200 OK\nMaksimum 10 log döner\nSıralama: en yeni önce", "priority": "ORTA"},

    {"id": "LOG-004", "module": "System Logs", "submodule": "Oluşturma",
     "scenario": "Yeni log kaydı oluşturma",
     "preconditions": "- Backend çalışıyor",
     "steps": "1. POST /api/v1/system-logs\n2. Body: {\"assetId\":\"DRONE-001\", \"eventType\":\"MODE_CHANGE\", \"message\":\"Mode changed to AUTO\"}",
     "expected": "200 OK\nLog kaydı oluşturulur\ncreatedAt otomatik set edilir", "priority": "ORTA"},
]

# ============================================================================
# TEST DATA — 11. FRONTEND UI
# ============================================================================
frontend_tests = [
    {"section": "25. GİRİŞ SAYFASI (LOGIN)", "id": "FE-001", "module": "Frontend", "submodule": "Login Page",
     "scenario": "Başarılı giriş ve yönlendirme",
     "preconditions": "- Frontend çalışıyor (localhost:5173)\n- Geçerli kullanıcı mevcut",
     "steps": "1. Login sayfasına git\n2. Doğru username/password gir\n3. Giriş yap butonuna tıkla",
     "expected": "JWT token localStorage'a kaydedilir\nKullanıcı Dashboard'a yönlendirilir", "priority": "KRİTİK"},

    {"id": "FE-002", "module": "Frontend", "submodule": "Login Page",
     "scenario": "Hatalı giriş uyarısı",
     "preconditions": "- Frontend çalışıyor",
     "steps": "1. Yanlış username/password gir\n2. Giriş yap butonuna tıkla",
     "expected": "Hata mesajı görünür\nSayfa yenilenmez\nToken oluşmaz", "priority": "KRİTİK"},

    {"id": "FE-003", "module": "Frontend", "submodule": "Login Page",
     "scenario": "Token olmadan korumalı sayfaya erişim",
     "preconditions": "- LocalStorage'da token yok",
     "steps": "1. Direkt /map URL'ine git",
     "expected": "/login sayfasına yönlendirilir", "priority": "YÜKSEK"},

    {"section": "26. TAKTİK HARİTA", "id": "FE-004", "module": "Frontend", "submodule": "Map View",
     "scenario": "Harita üzerinde asset ikonlarının gösterimi",
     "preconditions": "- Giriş yapılmış\n- Simülasyon çalışıyor",
     "steps": "1. Harita sayfasına git\n2. Asset ikonlarının konumlarını gözlemle",
     "expected": "Tüm aktif asset'ler harita üzerinde doğru konumda\nTip bazlı ikon farklılığı (drone/vehicle/soldier/sensor)", "priority": "KRİTİK"},

    {"id": "FE-005", "module": "Frontend", "submodule": "Map View",
     "scenario": "Harita üzerinde gerçek zamanlı hareket",
     "preconditions": "- Simülasyon çalışıyor\n- Deploy edilmiş asset var",
     "steps": "1. Haritayı izle\n2. Asset ikonlarının hareket ettiğini gözlemle",
     "expected": "İkonlar düzenli aralıklarla güncelleniyor\nHareket akıcı ve gözle takip edilebilir", "priority": "KRİTİK"},

    {"id": "FE-006", "module": "Frontend", "submodule": "Map View",
     "scenario": "Taktik bölgelerin haritada gösterimi",
     "preconditions": "- Tactical zones tanımlı",
     "steps": "1. Harita üzerinde zone çemberlerini gözlemle",
     "expected": "Her zone çember/polygon olarak gösterilir\nİsim etiketi görünür", "priority": "YÜKSEK"},

    {"id": "FE-007", "module": "Frontend", "submodule": "Map View",
     "scenario": "Asset'e tıklayınca popup/detay bilgisi",
     "preconditions": "- Harita üzerinde asset'ler var",
     "steps": "1. Bir asset ikonuna tıkla\n2. Popup/panel içeriğini kontrol et",
     "expected": "Asset adı, tipi, batarya, hız, konum bilgileri görünür", "priority": "YÜKSEK"},

    {"section": "27. TELEMETRİ DASHBOARD", "id": "FE-008", "module": "Frontend", "submodule": "Telemetry",
     "scenario": "Telemetri dashboard'unda veri gösterimi",
     "preconditions": "- Giriş yapılmış\n- Telemetri verisi var",
     "steps": "1. Telemetri sayfasına git\n2. Veri tablolarını/kartlarını kontrol et",
     "expected": "Tüm asset'ler için güncel telemetri gösterilir\nBatarya, hız, konum bilgileri doğru", "priority": "KRİTİK"},

    {"section": "28. ALARM PANELİ", "id": "FE-009", "module": "Frontend", "submodule": "Alerts Page",
     "scenario": "Alarm listesi gösterimi",
     "preconditions": "- Alarm kayıtları var",
     "steps": "1. Alarmlar sayfasına git\n2. Alarm listesini kontrol et",
     "expected": "Alarmlar severity'ye göre renklendirilmiş\nSıralama: en yeni önce", "priority": "KRİTİK"},

    {"id": "FE-010", "module": "Frontend", "submodule": "Alerts Page",
     "scenario": "Alarm onaylama (acknowledge) UI'dan",
     "preconditions": "- Onaylanmamış alarm var",
     "steps": "1. Alarmlar sayfasında ilgili alarmı bul\n2. Onayla butonuna tıkla",
     "expected": "Alarm onaylanmış olarak işaretlenir\nUI güncel durumu yansıtır", "priority": "YÜKSEK"},

    {"section": "29. AI SOHBET ARAYÜZÜ", "id": "FE-011", "module": "Frontend", "submodule": "AI Chat",
     "scenario": "AI sohbet oturumu oluşturma ve mesaj gönderme",
     "preconditions": "- Giriş yapılmış\n- AI Service çalışıyor",
     "steps": "1. AI Chat sayfasına git\n2. Yeni oturum oluştur\n3. Mesaj yaz ve gönder\n4. AI yanıtını bekle",
     "expected": "Mesaj gönderilir\nYükleniyor animasyonu gösterilir\nAI yanıtı görüntülenir\ncited_rule (varsa) gösterilir", "priority": "KRİTİK"},

    {"id": "FE-012", "module": "Frontend", "submodule": "AI Chat",
     "scenario": "Sohbet geçmişi yükleme",
     "preconditions": "- Geçmiş mesajları olan oturum var",
     "steps": "1. Mevcut oturuma tıkla\n2. Geçmiş mesajları kontrol et",
     "expected": "Tüm geçmiş mesajlar kronolojik sırada yüklenir\nuser/assistant rolleri doğru gösterilir", "priority": "YÜKSEK"},

    {"section": "30. ADMIN PANELİ UI", "id": "FE-013", "module": "Frontend", "submodule": "Admin Dashboard",
     "scenario": "Admin dashboard istatistikleri",
     "preconditions": "- KARARGAH_ADMIN ile giriş",
     "steps": "1. Admin Dashboard'a git\n2. İstatistik kartlarını kontrol et",
     "expected": "Aktif kullanıcı, toplam asset, aktif görev, toplam zone sayıları\nServis durumları (backend, DB, AI, MQTT)", "priority": "KRİTİK"},

    {"id": "FE-014", "module": "Frontend", "submodule": "Admin - User Mgmt",
     "scenario": "Admin panelinden kullanıcı oluşturma",
     "preconditions": "- KARARGAH_ADMIN ile giriş\n- Kullanıcı yönetim sayfası açık",
     "steps": "1. 'Yeni Kullanıcı' formunu doldur\n2. Rol seç\n3. Zone ata (HAREKAT_SUBAYI için)\n4. Kaydet",
     "expected": "Kullanıcı başarıyla oluşturulur\nListe güncellenir\nBaşarı bildirimi gösterilir", "priority": "KRİTİK"},

    {"id": "FE-015", "module": "Frontend", "submodule": "Admin - Asset Registry",
     "scenario": "Admin panelinden asset oluşturma",
     "preconditions": "- KARARGAH_ADMIN ile giriş\n- Asset kayıt sayfası açık",
     "steps": "1. Asset kayıt formunu doldur (id, tip, koordinatlar)\n2. Kaydet\n3. Listenin güncellendiğini doğrula",
     "expected": "Asset oluşturulur\nİlk telemetri kaydı otomatik oluşur\nListe güncellenir", "priority": "KRİTİK"},

    {"id": "FE-016", "module": "Frontend", "submodule": "Admin - System Health",
     "scenario": "Sistem sağlığı sayfası",
     "preconditions": "- KARARGAH_ADMIN ile giriş",
     "steps": "1. Sistem Sağlığı sayfasına git\n2. Her servisin durumunu kontrol et",
     "expected": "Backend: ONLINE\nDatabase: ONLINE/OFFLINE\nAI Service: ONLINE/OFFLINE\nMQTT: ONLINE/OFFLINE\nDocker: durumu", "priority": "YÜKSEK"},

    {"section": "31. NAVIGASYON & RESPONSIVE", "id": "FE-017", "module": "Frontend", "submodule": "Navigation",
     "scenario": "Sidebar navigasyon tüm sayfalara erişim",
     "preconditions": "- Giriş yapılmış",
     "steps": "1. Sidebar'dan her menü öğesine sırayla tıkla\n2. Her sayfanın yüklendiğini doğrula",
     "expected": "Tüm sayfalar sorunsuz yüklenir\nAktif menü öğesi vurgulanır", "priority": "YÜKSEK"},

    {"id": "FE-018", "module": "Frontend", "submodule": "Logout",
     "scenario": "Çıkış yapma",
     "preconditions": "- Giriş yapılmış",
     "steps": "1. Çıkış butonuna tıkla\n2. LocalStorage'ı kontrol et",
     "expected": "Token localStorage'dan silinir\nLogin sayfasına yönlendirilir", "priority": "KRİTİK"},
]

# ============================================================================
# TEST DATA — 12. SIMULATION & INTEGRATION
# ============================================================================
sim_tests = [
    {"section": "32. SİMÜLASYON MOTORU", "id": "SIM-001", "module": "Simulation", "submodule": "Engine",
     "scenario": "Simülasyon motorunun telemetri üretmesi",
     "preconditions": "- Backend çalışıyor\n- SimulationEngine aktif\n- Asset'ler mevcut",
     "steps": "1. Backend'i başlat\n2. DB'deki telemetry kayıtlarını izle\n3. Yeni kayıtların eklendiğini doğrula",
     "expected": "Her 3 saniyede tüm asset'ler için telemetri güncellenir\nKonum, hız, batarya değerleri değişir", "priority": "KRİTİK"},

    {"id": "SIM-002", "module": "Simulation", "submodule": "Movement",
     "scenario": "Deploy edilen asset'in hedef zone'a hareketi",
     "preconditions": "- Asset ve zone mevcut\n- Simülasyon çalışıyor",
     "steps": "1. Asset'i zone'a deploy et\n2. Haritada hareketi izle\n3. Varışı bekle",
     "expected": "Asset hedef zone'a doğru hareket eder\npayload.status: EN_ROUTE → IN_MISSION\ndistance_remaining_m azalır", "priority": "KRİTİK"},

    {"id": "SIM-003", "module": "Simulation", "submodule": "Scatter",
     "scenario": "Zone içindeki birimlerin saçılmış konumlanması",
     "preconditions": "- Birden fazla asset aynı zone'a deploy edilmiş",
     "steps": "1. Birden fazla asset'i aynı zone'a konuşlandır\n2. Haritada konumları gözlemle",
     "expected": "Asset'ler aynı noktada üst üste binmez\nZone yarıçapı içinde saçılmış konumlanır", "priority": "YÜKSEK"},

    {"section": "33. MQTT ENTEGRASYONU", "id": "SIM-004", "module": "Integration", "submodule": "MQTT",
     "scenario": "MQTT broker bağlantısı",
     "preconditions": "- MQTT broker çalışıyor (localhost:1883)",
     "steps": "1. Backend loglarını kontrol et\n2. \"MQTT Connection Successful\" mesajını doğrula",
     "expected": "Backend MQTT'ye bağlanır\nc4isr/+/telemetry topic'ine subscribe olur", "priority": "YÜKSEK"},

    {"id": "SIM-005", "module": "Integration", "submodule": "MQTT",
     "scenario": "MQTT üzerinden telemetri gönderme",
     "preconditions": "- MQTT broker ve backend çalışıyor",
     "steps": "1. MQTT client ile c4isr/drone/telemetry topic'ine mesaj yayınla\n2. Backend log'unu izle\n3. DB'de kayıt oluştuğunu doğrula",
     "expected": "Mesaj alınır\nTelemetry tablosuna kaydedilir", "priority": "YÜKSEK"},

    {"section": "34. SWAGGER UI / OPENAPI", "id": "SIM-006", "module": "Integration", "submodule": "Swagger",
     "scenario": "Swagger UI erişilebilirliği",
     "preconditions": "- Backend çalışıyor",
     "steps": "1. http://localhost:8080/swagger-ui/index.html adresine git",
     "expected": "Swagger UI yüklenir\nSINAPS C4ISR API başlığı görünür\nTüm endpoint'ler listelenir", "priority": "KRİTİK"},

    {"id": "SIM-007", "module": "Integration", "submodule": "Swagger",
     "scenario": "Swagger UI ile JWT Bearer token kullanımı",
     "preconditions": "- Swagger UI açık\n- Geçerli JWT token mevcut",
     "steps": "1. Authorize butonuna tıkla\n2. JWT token'ı yapıştır\n3. Authorize tıkla\n4. Korumalı endpoint'i \"Try it out\" ile test et",
     "expected": "Token kabul edilir\nKorumalı endpoint'ler 200 döner", "priority": "KRİTİK"},

    {"section": "35. PERFORMANS & DAYANIKLILIK", "id": "SIM-008", "module": "Integration", "submodule": "Performance",
     "scenario": "Yüksek sayıda asset ile simülasyon",
     "preconditions": "- 20+ asset DB'de mevcut\n- Simülasyon çalışıyor",
     "steps": "1. Backend'in CPU/memory kullanımını izle\n2. Harita üzerinde tüm asset'lerin gösterildiğini doğrula\n3. 10 dakika boyunca stabil çalıştığını doğrula",
     "expected": "Sistem stabil çalışır\nTelemetri pollingi aksamadan devam eder\n%100 CPU'ya ulaşmaz", "priority": "YÜKSEK"},

    {"id": "SIM-009", "module": "Integration", "submodule": "DB",
     "scenario": "Backend → PostgreSQL bağlantı testi",
     "preconditions": "- PostgreSQL çalışıyor (port 5433)",
     "steps": "1. Backend'i başlat\n2. HikariCP bağlantı loglarını kontrol et\n3. Herhangi bir GET endpoint'ini çağır",
     "expected": "Bağlantı başarılı\nHikariPool-1 initialized logu\nVeri başarılı şekilde çekilir", "priority": "KRİTİK"},

    {"id": "SIM-010", "module": "Integration", "submodule": "CORS",
     "scenario": "Frontend → Backend CORS politikası",
     "preconditions": "- Frontend (5173) ve Backend (8080) çalışıyor",
     "steps": "1. Frontend'den herhangi bir API çağrısı yap\n2. Browser console'da CORS hatası olup olmadığını kontrol et",
     "expected": "CORS hatası yok\nAccess-Control-Allow-Origin header'ı doğru\nPreflight (OPTIONS) istekleri başarılı", "priority": "KRİTİK"},
]


# ============================================================================
# GENERATE SHEETS
# ============================================================================
# Remove default sheet
del wb["Sheet"]

create_sheet("1-Auth & Security", auth_tests)
create_sheet("2-Asset Management", asset_tests)
create_sheet("3-Telemetry", telemetry_tests)
create_sheet("4-Tactical Zones", zone_tests)
create_sheet("5-Tasks & Deploy", task_tests)
create_sheet("6-Alerts & Analysis", alert_tests)
create_sheet("7-AI Service", ai_tests)
create_sheet("8-Chat Sessions", chat_tests)
create_sheet("9-Admin Panel", admin_tests)
create_sheet("10-System Logs", log_tests)
create_sheet("11-Frontend UI", frontend_tests)
create_sheet("12-Sim & Integration", sim_tests)

# ============================================================================
# SUMMARY SHEET (first tab)
# ============================================================================
summary = wb.create_sheet("ÖZET", 0)
summary_headers = ["Sheet", "Modül", "Toplam Test", "Başarılı", "Başarısız", "Atlanmış", "Oran (%)"]
for col_idx, h in enumerate(summary_headers, 1):
    cell = summary.cell(row=1, column=col_idx, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = CENTER
    cell.border = THIN_BORDER

summary.column_dimensions["A"].width = 25
summary.column_dimensions["B"].width = 22
summary.column_dimensions["C"].width = 14
summary.column_dimensions["D"].width = 12
summary.column_dimensions["E"].width = 12
summary.column_dimensions["F"].width = 12
summary.column_dimensions["G"].width = 12

sheet_data = [
    ("1-Auth & Security", "Kimlik Doğrulama & Güvenlik", len([t for t in auth_tests if t.get("id")])),
    ("2-Asset Management", "Varlık Yönetimi", len([t for t in asset_tests if t.get("id")])),
    ("3-Telemetry", "Telemetri", len([t for t in telemetry_tests if t.get("id")])),
    ("4-Tactical Zones", "Taktik Bölgeler", len([t for t in zone_tests if t.get("id")])),
    ("5-Tasks & Deploy", "Görevler & Konuşlandırma", len([t for t in task_tests if t.get("id")])),
    ("6-Alerts & Analysis", "Alarmlar & Analiz", len([t for t in alert_tests if t.get("id")])),
    ("7-AI Service", "Yapay Zeka Servisi", len([t for t in ai_tests if t.get("id")])),
    ("8-Chat Sessions", "Sohbet Oturumları", len([t for t in chat_tests if t.get("id")])),
    ("9-Admin Panel", "Admin Paneli", len([t for t in admin_tests if t.get("id")])),
    ("10-System Logs", "Sistem Logları", len([t for t in log_tests if t.get("id")])),
    ("11-Frontend UI", "Frontend Arayüzü", len([t for t in frontend_tests if t.get("id")])),
    ("12-Sim & Integration", "Simülasyon & Entegrasyon", len([t for t in sim_tests if t.get("id")])),
]

for row_idx, (sheet_name, module_name, total) in enumerate(sheet_data, 2):
    summary.cell(row=row_idx, column=1, value=sheet_name).border = THIN_BORDER
    summary.cell(row=row_idx, column=2, value=module_name).border = THIN_BORDER
    summary.cell(row=row_idx, column=3, value=total).border = THIN_BORDER
    summary.cell(row=row_idx, column=3).alignment = CENTER
    for c in range(4, 8):
        cell = summary.cell(row=row_idx, column=c, value=0)
        cell.border = THIN_BORDER
        cell.alignment = CENTER

# Total row
total_row = len(sheet_data) + 2
summary.cell(row=total_row, column=1, value="TOPLAM").font = Font(bold=True)
summary.cell(row=total_row, column=1).border = THIN_BORDER
summary.cell(row=total_row, column=2, value="").border = THIN_BORDER
for col in range(3, 7):
    cell = summary.cell(row=total_row, column=col)
    cell.value = f"=SUM({get_column_letter(col)}2:{get_column_letter(col)}{total_row-1})"
    cell.font = Font(bold=True)
    cell.border = THIN_BORDER
    cell.alignment = CENTER

# Pass rate formula
cell = summary.cell(row=total_row, column=7)
cell.value = f'=IF(C{total_row}>0, ROUND(D{total_row}/C{total_row}*100, 1), 0)'
cell.font = Font(bold=True)
cell.border = THIN_BORDER
cell.alignment = CENTER

# Metadata
summary.cell(row=total_row + 2, column=1, value="Proje:").font = Font(bold=True)
summary.cell(row=total_row + 2, column=2, value="SINAPS C4ISR Prototype")
summary.cell(row=total_row + 3, column=1, value="Oluşturulma:").font = Font(bold=True)
summary.cell(row=total_row + 3, column=2, value=datetime.datetime.now().strftime("%Y-%m-%d %H:%M"))
summary.cell(row=total_row + 4, column=1, value="Sürüm:").font = Font(bold=True)
summary.cell(row=total_row + 4, column=2, value="1.0")

summary.freeze_panes = "A2"

# ============================================================================
# SAVE
# ============================================================================
output_path = r"c:\Users\MCB\Desktop\Yazılım\Github\AI-Supported-Lightweight-C4ISR-Prototype-SINAPS\SINAPS_Test_Senaryolari.xlsx"
wb.save(output_path)
print(f"[OK] Test senaryolari Excel dosyasi olusturuldu: {output_path}")
print(f"   Toplam sheet: {len(wb.sheetnames)}")
total_tests = sum(s[2] for s in sheet_data)
print(f"   Toplam test senaryosu: {total_tests}")
